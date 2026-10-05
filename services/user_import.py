"""Bounded, non-mutating account spreadsheet validation. Passwords never enter preview output."""
from __future__ import annotations

import hashlib
from io import BytesIO
from itertools import islice
import json
import re
from types import SimpleNamespace

from services.account_management import validate_create_user
from services.company_config import company_from_settings
from services.work_order_import import validate_xlsx_archive, XlsxArchiveLimits

FIELDS = ('user_id', 'name', 'role', 'team', 'line_scope', 'password')
MAX_BYTES = 1024 * 1024
MAX_ROWS = 200


def read_rows(content: bytes) -> list:
    if len(content) > MAX_BYTES:
        raise ValueError('User file exceeds 1 MB limit')
    error = validate_xlsx_archive(content, XlsxArchiveLimits(1000, 16 * 1024 * 1024, 4 * 1024 * 1024, 200))
    if error:
        raise ValueError(error)
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet.max_column and sheet.max_column > len(FIELDS):
            raise ValueError('User import supports only the six template columns')
        result = []
        for row in islice(sheet.iter_rows(max_col=len(FIELDS)), MAX_ROWS + 2):
            if any(cell.data_type == 'f' for cell in row):
                raise ValueError('Formulas are not supported in user import')
            result.append(tuple(cell.value for cell in row))
        if len(result) > MAX_ROWS + 1:
            raise ValueError('User import supports at most 200 data rows')
        if not result:
            raise ValueError('User file is empty')
        return result
    finally:
        workbook.close()


def prepare(rows: list, accounts: dict, settings: dict) -> tuple[list[dict], list[dict]]:
    header = [str(value or '').strip() for value in rows[0]]
    populated = [value for value in header if value]
    if (set(populated) - set(FIELDS) or len(populated) != len(set(populated))
            or not {'user_id', 'role', 'line_scope', 'password'} <= set(populated)):
        raise ValueError('Use template headers: user_id, name, role, team, line_scope, password')
    lines = {line.id for line in company_from_settings(settings).lines}
    previews, accepted = [], []
    seen = set(accounts)
    for index, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        values = {key: str(row[col] if col < len(row) and row[col] is not None else '')
                  for col, key in enumerate(header) if key}
        values = {key: value if key == 'password' else value.strip() for key, value in values.items()}
        public = {key: value for key, value in values.items() if key != 'password'}
        preview: dict = {'row': index, 'values': public, 'error': '', 'warnings': ['首次登入必須更換密碼']}
        previews.append(preview)
        try:
            limits = dict(user_id=128, name=255, role=32, team=128, line_scope=2000, password=256)
            if any(len(value) > limits[key] for key, value in values.items()):
                raise ValueError('One or more fields exceeds the length limit')
            # Do not echo arbitrary input: it may include a password pasted into the wrong column.
            user_id = values['user_id']
            if not re.fullmatch(r'[A-Za-z0-9_-]{1,128}', user_id):
                raise ValueError('user_id must use ASCII letters, digits, dash or underscore')
            if user_id in seen:
                raise ValueError('Account already exists or is repeated in this file')
            if values['role'] not in {'operator', 'maintenance', 'supervisor', 'admin'}:
                raise ValueError('role must be operator, maintenance, supervisor or admin')
            scope = [s.strip() for s in values['line_scope'].split(',') if s.strip()]
            if not scope or ('*' in scope and scope != ['*']) or any(s != '*' and s not in lines for s in scope):
                raise ValueError('line_scope must contain configured line IDs, or * alone')
            req = SimpleNamespace(**{'name': '', 'team': '', **values, 'line_scope': list(dict.fromkeys(scope))})
            error = validate_create_user(req, accounts, valid_roles={'operator', 'maintenance', 'supervisor', 'admin'},
                                         default_password=lambda: '', missing_password_error=lambda: None)
            if error:
                raise ValueError(error)
            if values['role'] == 'admin' or scope == ['*']:
                preview['warnings'].append('請確認管理員角色或全部產線權限符合需要')
            public['line_scope'] = ','.join(req.line_scope)
            seen.add(user_id)
            accepted.append({'row': index, 'request': req})
        except ValueError as exc:
            preview['error'] = str(exc)
    return previews, accepted


def token(content: bytes, actor: str, accounts: dict, settings: dict) -> str:
    versions = {key: {field: value.get(field) for field in ('updated_at', 'role', 'active', 'credential_epoch', 'name', 'team', 'line_scope')}
                for key, value in accounts.items()}
    payload = {'file': hashlib.sha256(content).hexdigest(), 'actor': actor, 'accounts': versions, 'settings': settings}
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
